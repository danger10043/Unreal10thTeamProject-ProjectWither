from pathlib import Path
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def link_heading_text_to_toc(paragraph) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "TOC")
    hyperlink.set(qn("w:history"), "1")
    runs = list(paragraph._p.findall(qn("w:r")))
    if not runs:
        return
    insert_at = paragraph._p.index(runs[0])
    for run in runs:
        paragraph._p.remove(run)
        hyperlink.append(run)
    paragraph._p.insert(insert_at, hyperlink)


def main(input_path: Path, output_path: Path) -> None:
    doc = Document(input_path)

    if "TOCTitleCustom" not in [style.name for style in doc.styles]:
        style = doc.styles.add_style("TOCTitleCustom", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = doc.styles["Heading 1"].font.name
        style.font.size = Pt(20)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(14)
        style.paragraph_format.page_break_before = True

    first_heading = next(p for p in doc.paragraphs if p.style.name == "Heading 1")

    toc_title = OxmlElement("w:p")
    first_heading._p.addprevious(toc_title)
    from docx.text.paragraph import Paragraph
    toc_title_paragraph = Paragraph(toc_title, first_heading._parent)
    toc_title_paragraph.style = "TOCTitleCustom"
    toc_title_paragraph.add_run("목차")
    add_bookmark(toc_title_paragraph, "TOC", 10000)

    toc_placeholder = OxmlElement("w:p")
    first_heading._p.addprevious(toc_placeholder)
    placeholder_paragraph = Paragraph(toc_placeholder, first_heading._parent)
    placeholder_paragraph.style = "Normal"
    placeholder_paragraph.add_run("[[TOC]]")

    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            link_heading_text_to_toc(paragraph)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
